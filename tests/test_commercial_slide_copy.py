from pathlib import Path

from pptx import Presentation

TEMPLATE = Path("experiments/pptx_template/input/template.pptx")


def _shape(slide, name: str):
    return next(shape for shape in slide.shapes if shape.name == name)


def test_slide04_uses_generic_installation_commitments() -> None:
    presentation = Presentation(TEMPLATE)
    checklist = _shape(presentation.slides[3], "object 19").text

    assert "Protección y seguridad durante toda la instalación" in checklist
    assert "Trabajo profesional y cuidado de cada detalle" in checklist
    assert "Remates exteriores completamente terminados" in checklist
    assert "excepto remates de pintura y su preparación" in checklist
    assert "peana" not in checklist.lower()
    assert "execpto" not in checklist.lower()


def test_slide07_removes_advantages_box_and_keeps_clean_closing() -> None:
    presentation = Presentation(TEMPLATE)
    slide = presentation.slides[6]
    names = {shape.name for shape in slide.shapes}
    closing = _shape(slide, "object 12").text

    assert {"object 8", "object 9", "object 10", "object 11"}.isdisjoint(names)
    assert "Ventajas exclusivas" not in closing
    assert "Gracias por confiar en SmartVitra" in closing
    assert "vivienda más confortable y eficiente" in closing
