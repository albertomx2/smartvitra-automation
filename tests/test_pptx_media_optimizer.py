from __future__ import annotations

from pathlib import Path
from zipfile import ZipFile

from PIL import Image
from pptx import Presentation
from pptx.util import Inches

from backend.rendering.pptx.media_optimizer import (
    PptxMediaOptimizer,
)


def _create_noisy_image(
    path: Path,
    *,
    transparent: bool = False,
) -> None:
    mode = "RGBA" if transparent else "RGB"
    image = Image.effect_noise(
        (1600, 1200),
        80,
    ).convert(mode)

    if transparent:
        image.putalpha(180)

    image.save(
        path,
        format="PNG",
    )


def test_optimizer_reduces_photos_and_preserves_presentation_structure(
    tmp_path: Path,
) -> None:
    photo = tmp_path / "photo.png"
    _create_noisy_image(photo)

    presentation_path = tmp_path / "proposal.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    picture = slide.shapes.add_picture(
        str(photo),
        Inches(1),
        Inches(1),
        width=Inches(5),
        height=Inches(4),
    )
    picture.crop_left = 0.125
    picture.crop_right = 0.075
    textbox = slide.shapes.add_textbox(
        Inches(1),
        Inches(5.5),
        Inches(5),
        Inches(0.5),
    )
    textbox.text = "Texto comercial sin cambios"
    presentation.save(presentation_path)

    size_before = presentation_path.stat().st_size
    before = Presentation(presentation_path)
    before_picture = before.slides[0].shapes[0]
    before_geometry = (
        before_picture.left,
        before_picture.top,
        before_picture.width,
        before_picture.height,
        before_picture.crop_left,
        before_picture.crop_right,
    )

    PptxMediaOptimizer().optimize(
        presentation_path=presentation_path,
    )

    size_after = presentation_path.stat().st_size
    after = Presentation(presentation_path)
    after_picture = after.slides[0].shapes[0]
    after_geometry = (
        after_picture.left,
        after_picture.top,
        after_picture.width,
        after_picture.height,
        after_picture.crop_left,
        after_picture.crop_right,
    )

    assert size_after < size_before * 0.6
    assert len(after.slides) == len(before.slides) == 1
    assert after_geometry == before_geometry
    assert after.slides[0].shapes[1].text == "Texto comercial sin cambios"

    with ZipFile(presentation_path) as package:
        media = [name for name in package.namelist() if name.startswith("ppt/media/")]

    assert any(name.endswith(".jpeg") for name in media)
    assert not any(name.endswith(".png") for name in media)


def test_optimizer_keeps_transparent_png_unchanged(
    tmp_path: Path,
) -> None:
    image_path = tmp_path / "transparent.png"
    _create_noisy_image(
        image_path,
        transparent=True,
    )

    presentation_path = tmp_path / "transparent.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    slide.shapes.add_picture(
        str(image_path),
        Inches(1),
        Inches(1),
    )
    presentation.save(presentation_path)

    PptxMediaOptimizer().optimize(
        presentation_path=presentation_path,
    )

    with ZipFile(presentation_path) as package:
        media = [name for name in package.namelist() if name.startswith("ppt/media/")]

    assert any(name.endswith(".png") for name in media)
    assert not any(name.endswith(".jpeg") for name in media)
