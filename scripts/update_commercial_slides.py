from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_AUTO_SIZE, PP_ALIGN
from pptx.util import Inches, Pt

TEMPLATE = Path("experiments/pptx_template/input/template.pptx")


def _shape(slide, name: str):
    return next(shape for shape in slide.shapes if shape.name == name)


def _remove_shape(slide, name: str) -> None:
    shape = next((shape for shape in slide.shapes if shape.name == name), None)
    if shape is None:
        return
    parent = shape._element.getparent()
    if parent is None:
        raise RuntimeError(f"Shape {name} has no XML parent")
    parent.remove(shape._element)


def _set_checked_paragraph(paragraph, text: str) -> None:
    paragraph.clear()
    check = paragraph.add_run()
    check.text = "✔ "
    check.font.name = "Segoe UI Symbol"
    check.font.size = Pt(15.5)
    body = paragraph.add_run()
    body.text = text
    body.font.name = "Liberation Sans"
    body.font.size = Pt(13.5)


def _update_slide04(presentation: Presentation) -> None:
    slide = presentation.slides[3]
    shape = _shape(slide, "object 19")
    messages = [
        "Protección y seguridad durante toda la instalación",
        "Trabajo profesional y cuidado de cada detalle",
        "Remates exteriores completamente terminados",
        "Acabado final listo, excepto remates de pintura y su preparación",
    ]
    if len(shape.text_frame.paragraphs) != len(messages):
        raise RuntimeError("Unexpected slide 4 checklist structure")
    for paragraph, message in zip(shape.text_frame.paragraphs, messages, strict=True):
        _set_checked_paragraph(paragraph, message)


def _update_slide07(presentation: Presentation) -> None:
    slide = presentation.slides[6]

    for name in ("object 8", "object 9", "object 10", "object 11"):
        _remove_shape(slide, name)

    accent = _shape(slide, "object 7")
    accent.left = Inches(1.05)
    accent.top = Inches(5.42)
    accent.width = Inches(5.95)
    accent.height = Inches(0.025)

    closing = _shape(slide, "object 12")
    closing.left = Inches(0.82)
    closing.top = Inches(5.58)
    closing.width = Inches(6.45)
    closing.height = Inches(0.95)
    closing.text_frame.clear()
    closing.text_frame.word_wrap = True
    closing.text_frame.auto_size = MSO_AUTO_SIZE.NONE

    title = closing.text_frame.paragraphs[0]
    title.alignment = PP_ALIGN.CENTER
    title_run = title.add_run()
    title_run.text = "Gracias por confiar en SmartVitra"
    title_run.font.name = "Liberation Serif"
    title_run.font.size = Pt(16.5)
    title_run.font.bold = True
    title_run.font.color.rgb = RGBColor(0x00, 0x7F, 0xB5)

    detail = closing.text_frame.add_paragraph()
    detail.alignment = PP_ALIGN.CENTER
    detail_run = detail.add_run()
    detail_run.text = (
        "Nuestro compromiso es que disfrutes de una vivienda más confortable "
        "y eficiente durante muchos años."
    )
    detail_run.font.name = "Liberation Sans"
    detail_run.font.size = Pt(12.5)
    detail_run.font.italic = True
    detail_run.font.color.rgb = RGBColor(0x26, 0x24, 0x24)


def main() -> None:
    presentation = Presentation(TEMPLATE)
    _update_slide04(presentation)
    _update_slide07(presentation)
    presentation.save(TEMPLATE)


if __name__ == "__main__":
    main()
