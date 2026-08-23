from __future__ import annotations

import argparse
from pathlib import Path
from typing import Any

from pptx import Presentation as load_presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.presentation import Presentation
from pptx.util import Pt

REVIEW_URL = "https://g.page/r/CfoccmQ83BZCEBM/review"


def find_shape(
    slide: Any,
    name: str,
) -> Any:
    for shape in slide.shapes:
        if shape.name == name:
            return shape

    raise RuntimeError(f"Shape {name!r} not found")


def remove_shape_by_name(
    slide: Any,
    name: str,
) -> None:
    for shape in list(slide.shapes):
        if shape.name != name:
            continue

        element = shape._element
        parent = element.getparent()

        if parent is not None:
            parent.remove(element)


def fix_slide_2(
    prs: Presentation,
) -> None:
    slide = prs.slides[1]

    # Las tres filas disponen ahora de
    # más separación vertical.
    row_tops = {
        1: 2250000,
        2: 3750000,
        3: 5250000,
    }

    pairs = [
        (
            "object 4",
            "sv_s02_issue_1",
            1,
        ),
        (
            "object 6",
            "sv_s02_issue_2",
            1,
        ),
        (
            "object 8",
            "sv_s02_issue_3",
            2,
        ),
        (
            "object 10",
            "sv_s02_issue_4",
            2,
        ),
        (
            "object 12",
            "sv_s02_issue_5",
            3,
        ),
        (
            "object 14",
            "sv_s02_issue_6",
            3,
        ),
    ]

    for (
        arrow_name,
        text_name,
        row,
    ) in pairs:
        arrow = find_shape(
            slide,
            arrow_name,
        )

        text = find_shape(
            slide,
            text_name,
        )

        top = row_tops[row]

        text.top = top

        text.height = 1050000 if text_name == "sv_s02_issue_6" else 820000

        arrow.top = top + 85000


def fix_slide_3(
    prs: Presentation,
) -> None:
    slide = prs.slides[2]

    secondary = find_shape(
        slide,
        "sv_s03_main_benefit_secondary",
    )

    # La caja tenía sólo 520700 EMU
    # aunque existe mucho más hueco debajo.
    secondary.height = 950000

    if secondary.has_text_frame:
        secondary.text_frame.word_wrap = True


def replace_review_star_characters(
    shape: Any,
) -> None:
    if not shape.has_text_frame:
        return

    remaining = 5

    for paragraph in shape.text_frame.paragraphs:
        for run in paragraph.runs:
            if remaining <= 0:
                return

            if "★" not in run.text:
                continue

            result: list[str] = []

            for character in run.text:
                if character == "★" and remaining > 0:
                    # Dejamos hueco visual para
                    # las estrellas vectoriales.
                    result.append("\u00a0")
                    remaining -= 1
                else:
                    result.append(character)

            run.text = "".join(result)


def add_vector_stars(
    slide: Any,
    *,
    text_shape_name: str,
    group_name: str,
) -> None:
    text_shape = find_shape(
        slide,
        text_shape_name,
    )

    replace_review_star_characters(
        text_shape,
    )

    # Permite ejecutar el parche varias
    # veces sin duplicar estrellas.
    for index in range(
        1,
        6,
    ):
        remove_shape_by_name(
            slide,
            f"{group_name}_{index}",
        )

    star_size = Pt(8.5)
    gap = Pt(1.5)

    start_x = text_shape.left
    start_y = text_shape.top + Pt(1)

    for index in range(5):
        star = slide.shapes.add_shape(
            MSO_SHAPE.STAR_5_POINT,
            start_x + index * (star_size + gap),
            start_y,
            star_size,
            star_size,
        )

        star.name = f"{group_name}_" f"{index + 1}"

        star.fill.solid()

        star.fill.fore_color.rgb = RGBColor(
            0x1A,
            0x98,
            0xC8,
        )

        star.line.fill.background()


def add_review_hyperlink(
    slide: Any,
) -> None:
    shape = find_shape(
        slide,
        "object 22",
    )

    if not shape.has_text_frame:
        raise RuntimeError("Review URL shape has " "no text frame")

    run_count = 0

    # PowerPoint puede almacenar una URL
    # aparentemente continua dividida en
    # varios runs.
    #
    # Como esta caja contiene únicamente
    # la URL, hacemos clicable cada run
    # que tenga texto.
    for paragraph in shape.text_frame.paragraphs:
        for run in paragraph.runs:
            if not run.text:
                continue

            run.hyperlink.address = REVIEW_URL

            run_count += 1

    if run_count == 0:
        raise RuntimeError("Review URL shape contains " "no text runs")


def fix_slide_6(
    prs: Presentation,
) -> None:
    slide = prs.slides[5]

    review_shapes = [
        (
            "object 7",
            "sv_s06_rating_1",
        ),
        (
            "object 13",
            "sv_s06_rating_2",
        ),
        (
            "object 19",
            "sv_s06_rating_3",
        ),
    ]

    for (
        text_shape_name,
        group_name,
    ) in review_shapes:
        add_vector_stars(
            slide,
            text_shape_name=(text_shape_name),
            group_name=group_name,
        )

    add_review_hyperlink(
        slide,
    )


def patch(
    input_path: Path,
    output_path: Path,
) -> None:
    if not input_path.exists():
        raise FileNotFoundError(input_path)

    prs = load_presentation(str(input_path))

    if len(prs.slides) < 6:
        raise RuntimeError("Presentation has fewer " "than 6 slides")

    fix_slide_2(prs)
    fix_slide_3(prs)
    fix_slide_6(prs)

    output_path.parent.mkdir(
        parents=True,
        exist_ok=True,
    )

    prs.save(str(output_path))


def main() -> None:
    parser = argparse.ArgumentParser()

    parser.add_argument(
        "input",
        type=Path,
    )

    parser.add_argument(
        "--output",
        type=Path,
        default=None,
    )

    args = parser.parse_args()

    output_path = args.output if args.output is not None else args.input

    patch(
        args.input,
        output_path,
    )

    print()
    print("=" * 70)
    print("PRESENTACIÓN CORREGIDA")
    print("=" * 70)
    print(f"Archivo: {output_path}")
    print("Slide 2: distribución vertical corregida")
    print("Slide 3: beneficio secundario ampliado")
    print("Slide 6: estrellas vectoriales añadidas")
    print("Slide 6: enlace de Google configurado")


if __name__ == "__main__":
    main()
