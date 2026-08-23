from __future__ import annotations

import argparse
from pathlib import Path

from pptx import Presentation

SOURCE_SHAPE = "sv_s03_main_benefit"
TARGET_SHAPE = "sv_s03_main_benefit_secondary"


def find_shape(slide, name: str):
    for shape in slide.shapes:
        if shape.name == name:
            return shape

    raise RuntimeError(f"Shape {name!r} not found")


def copy_body_style(
    *,
    source_shape,
    target_shape,
) -> None:
    source_paragraphs = source_shape.text_frame.paragraphs

    if len(source_paragraphs) < 2:
        raise RuntimeError("Source main benefit does not " "contain paragraph 1")

    source_paragraph = source_paragraphs[1]

    if not source_paragraph.runs:
        raise RuntimeError("Source body paragraph has no runs")

    source_run = source_paragraph.runs[0]

    source_font = source_run.font

    target_paragraphs = target_shape.text_frame.paragraphs

    if not target_paragraphs:
        raise RuntimeError("Target secondary benefit " "has no paragraph")

    target_paragraph = target_paragraphs[0]

    if not target_paragraph.runs:
        target_run = target_paragraph.add_run()
    else:
        target_run = target_paragraph.runs[0]

    target_font = target_run.font

    target_font.name = source_font.name
    target_font.size = source_font.size
    target_font.bold = source_font.bold
    target_font.italic = source_font.italic

    if source_font.color.type is not None and source_font.color.rgb is not None:
        target_font.color.rgb = source_font.color.rgb

    target_paragraph.alignment = source_paragraph.alignment

    target_paragraph.level = source_paragraph.level

    target_shape.text_frame.word_wrap = True


def patch(
    path: Path,
) -> None:
    prs = Presentation(str(path))

    slide = prs.slides[2]

    source = find_shape(
        slide,
        SOURCE_SHAPE,
    )

    target = find_shape(
        slide,
        TARGET_SHAPE,
    )

    copy_body_style(
        source_shape=source,
        target_shape=target,
    )

    prs.save(str(path))

    print("Slide 3 secondary benefit " "style synchronized")
    print(f"Archivo: {path}")


def main() -> None:
    parser = argparse.ArgumentParser()

    parser.add_argument(
        "pptx",
        type=Path,
    )

    args = parser.parse_args()

    patch(args.pptx)


if __name__ == "__main__":
    main()
